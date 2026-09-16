"""Microsoft Word (.docx) parser. Requires python-docx."""
from __future__ import annotations

from pathlib import Path

from .base import BaseParser, ParsedDocument


class DocxParser(BaseParser):
    def parse(self, file_path: str) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "python-docx is required for .docx files: pip install python-docx"
            )

        doc = Document(file_path)
        sections: list[dict[str, str]] = []
        current_title = ""
        current_paragraphs: list[str] = []

        for para in doc.paragraphs:
            if para.style.name.startswith("Heading"):
                if current_paragraphs:
                    sections.append({"title": current_title, "content": "\n".join(current_paragraphs)})
                current_title = para.text
                current_paragraphs = []
            elif para.text.strip():
                current_paragraphs.append(para.text)

        if current_paragraphs:
            sections.append({"title": current_title, "content": "\n".join(current_paragraphs)})

        content = "\n\n".join(
            (f"## {s['title']}\n{s['content']}" if s["title"] else s["content"])
            for s in sections
        )

        metadata = self.extract_metadata(file_path)
        metadata["word_count"] = len(content.split())
        if sections:
            metadata["title"] = sections[0]["title"] or Path(file_path).stem

        return ParsedDocument(content=content, sections=sections, metadata=metadata)
